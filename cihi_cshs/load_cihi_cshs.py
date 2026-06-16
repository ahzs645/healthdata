#!/usr/bin/env python3
"""
CIHI Cost of a Standard Hospital Stay (CSHS) Loader
===================================================
Parses CIHI "Cost of a Standard Hospital Stay" Excel exports into SQLite. This is
the public per-stay hospital cost input for the NHHR cost-benefit analysis
(benefit line B1 -- Healthcare Cost Savings).

CSHS is published through an INTERACTIVE tool (cihi.ca), not an API, so this is a
LOADER, not a scraper: download the BC and/or Canada exports from
    https://www.cihi.ca/en/indicators/cost-of-a-standard-hospital-stay
drop the .xlsx files into ./source/, and run this to (re)build the database.
Any number of CIHI CSHS exports in ./source/ are merged; the BC export carries
province/region totals, the Canada export adds peer-group + urban/rural tags.

Usage:
    python load_cihi_cshs.py
    python load_cihi_cshs.py --db cihi_cshs.db --source-dir source
"""
from __future__ import annotations

import argparse
import glob
import os
import sqlite3
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    import sys
    sys.exit("openpyxl is required: pip install openpyxl")

# Source header -> normalized column. Headers absent in a given file stay NULL.
HEADER_MAP = {
    "Province/territory": "province",
    "Reporting level": "reporting_level",
    "Region": "region",
    "Place or organization": "place_or_org",
    "Corporation": "corporation",
    "Time scale": "time_scale",
    "Time frame": "time_frame",
    "Indicator": "indicator",
    "CSHS": "cshs_raw",
    "Unit of measure": "unit",
    "Confidence interval lower limit": "ci_lower",
    "Confidence interval upper limit": "ci_upper",
    "Performance comparison": "performance_comparison",
    "Performance trend": "performance_trend",
    "Urban or rural/remote": "urban_rural",
    "Hospital peer group": "hospital_peer_group",
    "Long-Term Care Facility Size": "ltc_size",
    "Trend note": "trend_note",
    "Refresh date": "refresh_date",
}
COLUMNS = list(dict.fromkeys(HEADER_MAP.values())) + ["cshs_value", "source_file", "loaded_at"]


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def to_value(raw):
    if raw is None:
        return None
    s = str(raw).replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None  # "Suppressed", "Not applicable", etc.


def init_db(conn):
    cols = ",\n            ".join(f"{c} TEXT" if c not in ("cshs_value",) else f"{c} REAL" for c in COLUMNS)
    conn.execute(f"CREATE TABLE IF NOT EXISTS cshs (\n            id INTEGER PRIMARY KEY AUTOINCREMENT,\n            {cols}\n        )")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_cshs_place ON cshs(place_or_org)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_cshs_year ON cshs(time_frame)")
    conn.commit()


def load_file(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {h: i for i, h in enumerate(header)}
    src = os.path.basename(path)
    out = []
    for row in rows:
        if row is None or all(c is None for c in row):
            continue
        rec = {c: None for c in COLUMNS}
        for src_h, col in HEADER_MAP.items():
            if src_h in idx:
                v = row[idx[src_h]]
                rec[col] = str(v) if v is not None else None
        rec["cshs_value"] = to_value(rec.get("cshs_raw"))
        rec["source_file"] = src
        out.append(rec)
    return out


def main():
    ap = argparse.ArgumentParser(description="Load CIHI CSHS Excel exports into SQLite.")
    ap.add_argument("--db", default="cihi_cshs.db")
    ap.add_argument("--source-dir", default="source")
    args = ap.parse_args()

    files = sorted(glob.glob(os.path.join(args.source_dir, "*.xlsx")))
    if not files:
        raise SystemExit(f"No .xlsx files in {args.source_dir}/. Download CSHS exports from cihi.ca first.")

    conn = sqlite3.connect(args.db)
    init_db(conn)
    conn.execute("DELETE FROM cshs")  # full rebuild from current source files
    ts = now()
    total = 0
    placeholders = ",".join("?" for _ in COLUMNS)
    with conn:
        for path in files:
            recs = load_file(path)
            for r in recs:
                r["loaded_at"] = ts
                conn.execute(f"INSERT INTO cshs ({','.join(COLUMNS)}) VALUES ({placeholders})",
                             [r[c] for c in COLUMNS])
            total += len(recs)
            print(f"  loaded {len(recs):>5} rows from {os.path.basename(path)}")

    print(f"\n{args.db}: {total} rows from {len(files)} file(s)")

    def latest(place, level):
        c = conn.execute(
            "SELECT time_frame, cshs_value FROM cshs WHERE place_or_org=? AND reporting_level=? "
            "AND cshs_value IS NOT NULL ORDER BY time_frame DESC LIMIT 1", (place, level)).fetchone()
        return c

    print("Headline (latest fiscal year):")
    for place, level in [("British Columbia", "Province/territory"),
                         ("Northern Health (B.C.)", "Health region"),
                         ("University Hospital of Northern British Columbia (B.C.)", "Facility")]:
        r = latest(place, level)
        if r:
            print(f"  {place}: ${r[1]:,.0f}  ({r[0]})")
    conn.close()


if __name__ == "__main__":
    main()
